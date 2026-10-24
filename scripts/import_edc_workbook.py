from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sqlite3
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter


COMMON_INDEX_COLUMNS = {
    "项目编号",
    "中心编号",
    "中心名称",
    "病例编号",
    "患者姓名缩写",
    "受试者筛选编号",
    "受试者姓名缩写",
    "访视名称",
    "行号",
}


def quote_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def clean_header(value: Any, column_index: int) -> str:
    if value is None or str(value).strip() == "":
        return f"column_{get_column_letter(column_index)}"
    text = re.sub(r"\s+", " ", str(value).strip())
    return text


def unique_headers(values: list[Any]) -> list[str]:
    seen: Counter[str] = Counter()
    headers: list[str] = []
    for idx, value in enumerate(values, start=1):
        base = clean_header(value, idx)
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}__{seen[base]}")
    return headers


def normalize_value(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    if value == "":
        return None
    return value


def value_type(value: Any) -> str:
    if value is None or value == "":
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "real"
    if isinstance(value, dt.datetime):
        return "datetime"
    if isinstance(value, dt.date):
        return "date"
    if isinstance(value, dt.time):
        return "time"
    return "text"


def infer_sql_type(values: list[Any]) -> str:
    observed = {value_type(v) for v in values if v is not None and v != ""}
    if not observed:
        return "TEXT"
    if observed <= {"integer", "boolean"}:
        return "INTEGER"
    if observed <= {"integer", "real", "boolean"}:
        return "REAL"
    return "TEXT"


def table_name_for(index: int) -> str:
    return f"sheet_{index:02d}"


def workbook_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def fetch_row(ws: Any, row_index: int, max_column: int) -> list[Any]:
    return [
        ws.cell(row=row_index, column=col_index).value
        for col_index in range(1, max_column + 1)
    ]


def row_has_data(values: list[Any]) -> bool:
    return any(value is not None and value != "" for value in values)


def create_metadata_tables(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE _source_workbook (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            source_path TEXT NOT NULL,
            source_filename TEXT NOT NULL,
            file_size_bytes INTEGER NOT NULL,
            sha256 TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            workbook_sheet_count INTEGER NOT NULL
        );

        CREATE TABLE _sheets (
            sheet_index INTEGER PRIMARY KEY,
            sheet_name TEXT NOT NULL,
            table_name TEXT NOT NULL UNIQUE,
            sheet_state TEXT NOT NULL,
            header_row INTEGER NOT NULL,
            max_row INTEGER NOT NULL,
            max_column INTEGER NOT NULL,
            data_rows_imported INTEGER NOT NULL,
            non_empty_cells INTEGER NOT NULL
        );

        CREATE TABLE _columns (
            sheet_index INTEGER NOT NULL,
            table_name TEXT NOT NULL,
            excel_column_index INTEGER NOT NULL,
            excel_column_letter TEXT NOT NULL,
            original_header TEXT NOT NULL,
            sqlite_column_name TEXT NOT NULL,
            sqlite_type TEXT NOT NULL,
            PRIMARY KEY (sheet_index, excel_column_index)
        );

        CREATE TABLE excel_cells (
            sheet_index INTEGER NOT NULL,
            sheet_name TEXT NOT NULL,
            table_name TEXT NOT NULL,
            row_index INTEGER NOT NULL,
            column_index INTEGER NOT NULL,
            cell_address TEXT NOT NULL,
            column_header TEXT,
            value_text TEXT,
            value_type TEXT NOT NULL,
            formula_text TEXT,
            number_format TEXT,
            PRIMARY KEY (sheet_index, row_index, column_index)
        );

        CREATE INDEX idx_excel_cells_sheet_row ON excel_cells(sheet_name, row_index);
        CREATE INDEX idx_excel_cells_header ON excel_cells(column_header);
        """
    )


def try_create_fts(conn: sqlite3.Connection) -> bool:
    try:
        conn.execute(
            """
            CREATE VIRTUAL TABLE _rows_fts USING fts5(
                sheet_name UNINDEXED,
                table_name UNINDEXED,
                source_row UNINDEXED,
                row_text
            )
            """
        )
        conn.execute(
            """
            CREATE VIRTUAL TABLE _cells_fts USING fts5(
                sheet_name UNINDEXED,
                cell_address UNINDEXED,
                column_header,
                value_text
            )
            """
        )
        return True
    except sqlite3.OperationalError:
        return False


def import_workbook(source: Path, output: Path) -> dict[str, Any]:
    if output.exists():
        output.unlink()

    wb_values = openpyxl.load_workbook(source, data_only=True, read_only=False)
    wb_formulas = openpyxl.load_workbook(source, data_only=False, read_only=False)

    conn = sqlite3.connect(output)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.execute("PRAGMA foreign_keys=ON")

    create_metadata_tables(conn)
    has_fts = try_create_fts(conn)

    conn.execute(
        """
        INSERT INTO _source_workbook (
            id, source_path, source_filename, file_size_bytes, sha256,
            imported_at, workbook_sheet_count
        ) VALUES (1, ?, ?, ?, ?, ?, ?)
        """,
        (
            str(source),
            source.name,
            source.stat().st_size,
            workbook_sha256(source),
            dt.datetime.now().isoformat(timespec="seconds"),
            len(wb_values.worksheets),
        ),
    )

    report_sheets: list[dict[str, Any]] = []
    total_rows = 0
    total_cells = 0

    for sheet_index, ws_values in enumerate(wb_values.worksheets, start=1):
        ws_formulas = wb_formulas[ws_values.title]
        table_name = table_name_for(sheet_index)
        max_row = ws_values.max_row
        max_column = ws_values.max_column
        header_row = 1
        raw_headers = fetch_row(ws_values, header_row, max_column)
        headers = unique_headers(raw_headers)

        data_rows: list[list[Any]] = []
        column_samples: list[list[Any]] = [[] for _ in range(max_column)]
        for row_index in range(header_row + 1, max_row + 1):
            values = fetch_row(ws_values, row_index, max_column)
            if not row_has_data(values):
                continue
            normalized = [normalize_value(value) for value in values]
            data_rows.append([row_index, *normalized])
            for col_index, value in enumerate(normalized):
                column_samples[col_index].append(value)

        column_types = [infer_sql_type(values) for values in column_samples]
        table_columns = [
            "__source_row INTEGER NOT NULL",
            *(
                f"{quote_ident(header)} {column_type}"
                for header, column_type in zip(headers, column_types)
            ),
        ]
        conn.execute(
            f"CREATE TABLE {quote_ident(table_name)} ({', '.join(table_columns)})"
        )
        conn.execute(
            f"CREATE INDEX {quote_ident('idx_' + table_name + '_source_row')} "
            f"ON {quote_ident(table_name)} (__source_row)"
        )

        placeholders = ", ".join(["?"] * (max_column + 1))
        insert_sql = (
            f"INSERT INTO {quote_ident(table_name)} "
            f"({', '.join([quote_ident('__source_row'), *(quote_ident(h) for h in headers)])}) "
            f"VALUES ({placeholders})"
        )
        if data_rows:
            conn.executemany(insert_sql, data_rows)

        for header in headers:
            if header in COMMON_INDEX_COLUMNS:
                idx_name = f"idx_{table_name}_{len(header)}_{abs(hash(header)) % 100000}"
                conn.execute(
                    f"CREATE INDEX {quote_ident(idx_name)} "
                    f"ON {quote_ident(table_name)} ({quote_ident(header)})"
                )

        for col_index, (header, column_type) in enumerate(
            zip(headers, column_types), start=1
        ):
            conn.execute(
                """
                INSERT INTO _columns (
                    sheet_index, table_name, excel_column_index, excel_column_letter,
                    original_header, sqlite_column_name, sqlite_type
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    sheet_index,
                    table_name,
                    col_index,
                    get_column_letter(col_index),
                    clean_header(raw_headers[col_index - 1], col_index),
                    header,
                    column_type,
                ),
            )

        non_empty_cells = 0
        cell_records = []
        fts_cell_records = []
        for row in ws_values.iter_rows(
            min_row=1, max_row=max_row, min_col=1, max_col=max_column
        ):
            for cell in row:
                raw_value = cell.value
                if raw_value is None or raw_value == "":
                    continue
                formula_value = ws_formulas[cell.coordinate].value
                formula_text = (
                    formula_value
                    if isinstance(formula_value, str) and formula_value.startswith("=")
                    else None
                )
                col_header = headers[cell.column - 1] if cell.column <= len(headers) else None
                normalized_value = normalize_value(raw_value)
                value_text = None if normalized_value is None else str(normalized_value)
                cell_records.append(
                    (
                        sheet_index,
                        ws_values.title,
                        table_name,
                        cell.row,
                        cell.column,
                        cell.coordinate,
                        col_header,
                        value_text,
                        value_type(raw_value),
                        formula_text,
                        cell.number_format,
                    )
                )
                fts_cell_records.append(
                    (ws_values.title, cell.coordinate, col_header, value_text)
                )
                non_empty_cells += 1

        if cell_records:
            conn.executemany(
                """
                INSERT INTO excel_cells (
                    sheet_index, sheet_name, table_name, row_index, column_index,
                    cell_address, column_header, value_text, value_type, formula_text,
                    number_format
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                cell_records,
            )
            if has_fts:
                conn.executemany(
                    """
                    INSERT INTO _cells_fts (
                        sheet_name, cell_address, column_header, value_text
                    ) VALUES (?, ?, ?, ?)
                    """,
                    fts_cell_records,
                )

        if has_fts and data_rows:
            fts_rows = []
            for imported_row in data_rows:
                source_row = imported_row[0]
                values = imported_row[1:]
                parts = []
                for header, value in zip(headers, values):
                    if value is not None and value != "":
                        parts.append(f"{header}: {value}")
                fts_rows.append(
                    (ws_values.title, table_name, source_row, " | ".join(parts))
                )
            conn.executemany(
                """
                INSERT INTO _rows_fts (
                    sheet_name, table_name, source_row, row_text
                ) VALUES (?, ?, ?, ?)
                """,
                fts_rows,
            )

        conn.execute(
            """
            INSERT INTO _sheets (
                sheet_index, sheet_name, table_name, sheet_state, header_row,
                max_row, max_column, data_rows_imported, non_empty_cells
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                sheet_index,
                ws_values.title,
                table_name,
                ws_values.sheet_state,
                header_row,
                max_row,
                max_column,
                len(data_rows),
                non_empty_cells,
            ),
        )

        report_sheets.append(
            {
                "sheet_index": sheet_index,
                "sheet_name": ws_values.title,
                "table_name": table_name,
                "max_row": max_row,
                "max_column": max_column,
                "data_rows_imported": len(data_rows),
                "non_empty_cells": non_empty_cells,
            }
        )
        total_rows += len(data_rows)
        total_cells += non_empty_cells

    conn.execute("CREATE VIEW v_sheet_catalog AS SELECT * FROM _sheets ORDER BY sheet_index")
    conn.execute(
        """
        CREATE VIEW v_all_rows AS
        SELECT sheet_name, table_name, source_row, row_text
        FROM _rows_fts
        """
        if has_fts
        else """
        CREATE VIEW v_all_rows AS
        SELECT sheet_name, table_name, row_index AS source_row,
               group_concat(column_header || ': ' || value_text, ' | ') AS row_text
        FROM excel_cells
        GROUP BY sheet_name, table_name, row_index
        """
    )
    conn.commit()
    conn.execute("PRAGMA optimize")
    conn.close()

    return {
        "source": str(source),
        "output": str(output),
        "sheet_count": len(report_sheets),
        "total_data_rows_imported": total_rows,
        "total_non_empty_cells": total_cells,
        "fts_enabled": has_fts,
        "sheets": report_sheets,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    report = import_workbook(args.source, args.output)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
